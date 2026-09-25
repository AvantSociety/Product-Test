#!/usr/bin/env python3
"""
Builds the answer key workbook for a generated collection, and, when a run's
results.json sits next to it, fills in what the app actually did and marks
each check PASS, FAIL or FINDING.

    python3 answer_key.py out/120                   # key only
    python3 answer_key.py out/120 --scale out/250 out/500 ...

PASS     the app did what a correct review requires
FAIL     the app did something wrong (a defect to fix)
FINDING  the app followed its own rules, but the result is one an attorney
         must know about (a limit of the approach rather than a bug)
"""

import argparse
import csv
import datetime as dt
import io
import json
import os
import re
from collections import Counter

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from expectations import screen, eml_view, dates_in, LABEL


def norm(t):
    return re.sub(r'\s+', ' ', t).strip()


def passage_hit(passage, citations):
    """True when an extracted citation contains the passage's first sentence.
    Whitespace is normalised because PDF text comes back with spacing collapsed."""
    head = norm(re.split(r'(?<=[.!?])\s', norm(passage))[0])
    return [c for c in citations if head in norm(c['excerpt'])]

NAVY = '1E3A5F'
FILL = {'PASS': 'D9EAD3', 'FAIL': 'F4CCCC', 'FINDING': 'FFF2CC', 'n/a': 'EEEEEE'}
HOLD_LABELS = {
    'needs_ocr': 'Scanned image with no text layer', 'empty': 'No readable text',
    'corrupt': 'Corrupted character encoding', 'unrelated': 'Does not appear to belong to this matter',
    'no_connection': 'Names no party or key term of this matter',
}
# Either hold is a correct catch for a document from another matter.
FOREIGN = {'Does not appear to belong to this matter', 'Names no party or key term of this matter'}


def us(d):
    return f'{d.month}/{d.day}/{d.year}'


def load(dirpath):
    m = json.load(open(os.path.join(dirpath, 'manifest.json')))
    rpath = os.path.join(dirpath, 'results.json')
    r = json.load(open(rpath)) if os.path.exists(rpath) else None
    return m, r


def app_text(d):
    """The text the app should end up holding for this document."""
    if d['fmt'] in ('pdf_scan', 'empty'):
        return ''
    return eml_view(d['text']) if d['fmt'] == 'eml' else d['text']


def expected(m):
    """Stage-by-stage expectations derived from the ground truth."""
    docs = [d for d in m['documents'] if d.get('integrity') != 'duplicate']
    crit = m['criteria']
    exp = {'docs': {}}
    for d in docs:
        cat, p, t = screen(app_text(d), crit)
        exp['docs'][d['file']] = {'screen': cat, 'party_hits': p, 'term_hits': t}
    exp['accepted'] = len(docs)
    exp['screen_counts'] = Counter(v['screen'] for v in exp['docs'].values())
    exp['pass1_holds'] = {d['file']: d['integrity'] for d in docs if d.get('integrity') in HOLD_LABELS}
    # Any other readable document that matches none of the criteria is held
    # for confirmation too (the newsletter, the text-message export).
    for d in docs:
        if d['file'] not in exp['pass1_holds'] and app_text(d).strip() and d['fmt'] != 'cp1252' \
                and exp['docs'][d['file']]['screen'] == 'none':
            exp['pass1_holds'][d['file']] = 'no_connection'
    return exp


# ---------------------------------------------------------------------------

def build_workbook(dirpath, scale_dirs):
    m, r = load(dirpath)
    exp = expected(m)
    wb = Workbook()
    ws = wb.active
    ws.title = 'Read me'
    checks = stage_checks(m, r, exp)
    write_readme(ws, m, r, checks)
    write_checks(wb.create_sheet('Stage checks'), checks)
    write_documents(wb.create_sheet('Documents'), m, r, exp)
    write_chronology(wb.create_sheet('Key dates'), m, r)
    write_passages(wb.create_sheet('Must-find passages'), m, r)
    if scale_dirs:
        write_scale(wb.create_sheet('Scale results'), [dirpath] + scale_dirs)
    out = os.path.join(dirpath, f'Answer Key - {m["size"]} documents.xlsx')
    wb.save(out)
    return out, checks


def status(ok, finding=False):
    return 'PASS' if ok else ('FINDING' if finding else 'FAIL')


def stage_checks(m, r, exp):
    """Each check: (stage, what to check, expected, observed, result, why it matters)."""
    C = []
    docs = [d for d in m['documents'] if d.get('integrity') != 'duplicate']
    dup = next(d for d in m['documents'] if d.get('integrity') == 'duplicate')
    add = lambda *a: C.append(a)
    R = r or {}
    obs_docs = R.get('documents', {})

    # ---- Stage 01 ----
    skipped = [n for n in R.get('uploadNotices', []) if 'already ingested' in n]
    wrongly = [x.split(' was already ingested')[0] for x in skipped if dup['file'] not in x]
    add('01 Ingest', 'Every distinct file enters the matter',
        f"{exp['accepted']} of {m['size']} (only the one exact duplicate is skipped)",
        (f"{R.get('accepted')} entered; wrongly skipped as duplicates: " + '; '.join(wrongly)) if wrongly else R.get('accepted'),
        status(R.get('accepted') == exp['accepted']) if r else '',
        'DEFECT FOUND if files are wrongly skipped: duplicates are detected by hashing the extracted text. Every scanned PDF and every empty file has the same (empty) text, so the second and later ones are discarded as "already ingested", and are never reviewed, logged or produced. Fix: hash the file bytes, not the extracted text.')
    add('01 Ingest', 'Exact duplicate skipped with a notice', f'"{dup["file"]}" skipped as already ingested',
        '; '.join(x for x in skipped if dup['file'] in x) if r else '', status(any(dup['file'] in x for x in skipped)) if r else '',
        'Duplicates inflate volume and put two Bates numbers on one document.')
    scan_notice = [n for n in R.get('uploadNotices', []) if 'scanned image' in n]
    scans_in = [d['file'] for d in docs if d['fmt'] == 'pdf_scan' and (not r or d['file'] in obs_docs)]
    add('01 Ingest', 'Scanned PDFs that entered the matter are flagged for OCR at upload', f'{len(scans_in)} notice(s)', len(scan_notice) if r else '',
        status(len(scan_notice) == len(scans_in)) if r else '', 'The attorney needs to know which files the tool cannot read.')
    if r:
        wrong_type = [n for n, d in ((d['file'], d) for d in docs)
                      if n in obs_docs and obs_docs[n]['type'] != d['type'] and d.get('planted')]
        add('01 Ingest', 'Document type labels (48 planted documents)', 'Label matches the true document type',
            f'{48 - len([w for w in wrong_type])} of 48 match; differs: ' + '; '.join(f'{w} -> {obs_docs[w]["type"]}' for w in wrong_type),
            status(not wrong_type, finding=True),
            'The label is a convenience for sorting; a wrong label changes nothing downstream. Differences are listed so the team knows how far to trust it.')
        pdfs = [d for d in docs if d['fmt'] == 'pdf' and d['file'] in obs_docs]
        exact = sum(1 for d in pdfs if obs_docs[d['file']]['pagesExact'])
        add('01 Ingest', 'PDF page counts exact (not estimated)', f'{len(pdfs)} of {len(pdfs)}', f'{exact} of {len(pdfs)}',
            status(exact == len(pdfs)), 'Page counts go on the production index.')

    # ---- Stage 01 screening ----
    counts = exp['screen_counts']
    obs_counts = {}
    for chip in R.get('chips', []):
        k, v = chip.rsplit(' ', 1) if ' ' in chip else (chip, 0)
        obs_counts[k.strip()] = int(v)
    if r:
        accepted_names = set(obs_docs)
        counts = Counter(v['screen'] for n, v in exp['docs'].items() if n in accepted_names)
    exp_str = ', '.join(f'{LABEL[k]} {counts.get(k, 0)}' for k in ['strong', 'possible', 'out_of_period', 'none'])
    obs_str = ', '.join(f'{k} {v}' for k, v in obs_counts.items() if k != 'ALL')
    same = all(obs_counts.get(LABEL[k], 0) == counts.get(k, 0) for k in LABEL)
    add('01 Screening', 'Screening totals (for the documents the app accepted)', exp_str, obs_str if r else '',
        status(same) if r else '', 'Totals from an independent re-implementation of the documented rule.')
    if r:
        rows = R.get('rows', {})
        diff = [f'{n}: expected {LABEL[e["screen"]]}, app {rows.get(n, {}).get("screen")}'
                for n, e in exp['docs'].items() if n in rows and rows[n].get('screen') != LABEL[e['screen']]]
        add('01 Screening', 'Each document screened as the rule requires', 'All match', '; '.join(diff) or 'All match',
            status(not diff), 'Any difference is a screening defect.')
    for f, want, why in [
        ('2025-02-03 Carranza to Haldane - L3 supply trunks.eml', 'possible',
         'FINDING: the most important document in the case screens only POSSIBLE because it uses none of the key terms. Keyword screening is a queue, not a verdict.'),
        ('2022-08-09 Prequalification questionnaire - Tallis Belle Meade.pdf', 'out_of_period', 'Out-of-period documents are flagged, not excluded.'),
        ('Hollis v. Crane Property - demand letter re security deposit.docx', 'none', "Another client's document must not screen as relevant."),
        ('2025-02-18 Text messages - Carranza to Haldane.txt', 'none',
         'FINDING: a responsive text message screens NO MATCH ("Luis", "duct"). "Select relevant" leaves it out; the attorney must add it by hand.'),
        ('2024-11-08 Change Order 14 - MRI Suite Redesign (signed scan).pdf', 'none',
         'FINDING: a scan has no text, so it screens NO MATCH and "Select relevant" drops it before the Stage 03 OCR warning can fire. The signed Change Order 14 would silently not be produced.'),
    ]:
        got = R.get('rows', {}).get(f, {}).get('screen') if r else ''
        finding = why.startswith('FINDING')
        ok = got == LABEL[want]
        add('01 Screening', f'"{f}" screens {LABEL[want]}', LABEL[want], got,
            ('FINDING' if (ok and finding) else status(ok)) if r else '', why)

    # ---- Stage 02 ----
    wh = [d for d in docs if d['designation'] == 'withhold']
    rd = [d for d in docs if d['designation'] == 'redact']
    add('02 Designate', 'Withhold (privileged)', f'{len(wh)}: ' + '; '.join(f'{d["file"]} [{d["basis"]}]' for d in wh),
        R.get('selectionSummary', ''), status(R.get('selectionSummary', '').endswith(f'{len(wh)} withheld')) if r else '',
        'Three attorney-client emails and one work-product memo.')
    add('02 Designate', 'Redact (personal and financial identifiers)', f'{len(rd)}: ' + '; '.join(d['file'] for d in rd),
        ', '.join(n for n, v in (R.get('privilege') or {}).items() if v.get('status') == 'redact'),
        status(len([1 for v in (R.get('privilege') or {}).values() if v.get('status') == 'redact']) == len(rd)) if r else '',
        'Payroll SSNs and bank wire details. Basis "Other" with a description.')
    for d in docs:
        if d.get('judgment', '').startswith('TRAP'):
            add('02 Designate', f'TRAP - produce, do not withhold: {d["file"]}', 'Produce', 'Produce (attorney call; the tool does not decide privilege)', 'n/a', d['judgment'])

    # ---- Stage 03 ----
    holds_exp = exp['pass1_holds']
    p1 = {e['name']: e['defects'] for e in (R.get('pass1') or {}).get('exceptions', [])}
    for f, code in holds_exp.items():
        accepted = f in obs_docs if r else True
        got = '; '.join(p1.get(f, [])) if r else ''
        ok = (bool(FOREIGN & set(p1.get(f, []))) if code in ('unrelated', 'no_connection')
              else HOLD_LABELS[code] in p1.get(f, []))
        note = ''
        if code == 'no_connection':
            note = 'Readable, but matches none of the screening criteria. Held until the attorney confirms it belongs; released with one click if it does.'
        if code == 'unrelated' and r and not ok:
            note = ('DEFECT FOUND: this letter comes from our own law firm (Mabry & Colquitt) for another client. It shares the firm name, "LLP", "client" and "attorney" with our privileged emails, so the cohesion test sees it as related. '
                    'A misfiled document from the same firm is the most likely real-world intruder. Screening (NO MATCH) is the only thing that catches it.')
        if r and not accepted:
            note = 'Never reached Stage 03: it was skipped at upload as a false duplicate (see Stage 01).'
        res = ('n/a' if not accepted else status(ok)) if r else ''
        add('03 Readiness (pass 1: all selected)', f'Held: {f}',
            'Held as not belonging to this matter (either test)' if code in ('unrelated', 'no_connection') else HOLD_LABELS[code], got or ('not in matter' if not accepted else 'not held'),
            res, note or 'Blocking defect; the document cannot be produced as it stands.')
    extra = [f'{n} ({"; ".join(v)})' for n, v in p1.items() if n not in holds_exp]
    add('03 Readiness (pass 1: all selected)', 'Nothing else held back', 'None', '; '.join(extra) or 'None',
        status(not extra, finding=True) if r else '',
        'Anything listed here was held by the cohesion test. Review whether it truly belongs to another matter.')
    add('03 Readiness (pass 1: all selected)', 'Overall state', 'CURE REQUIRED', (R.get('pass1') or {}).get('state'),
        status((R.get('pass1') or {}).get('state') == 'CURE REQUIRED') if r else '', 'Under 25% unreadable, so cure rather than re-collect.')
    before = {e['name']: e['defects'] for e in (R.get('pass2BeforeConfirm') or R.get('pass2') or {}).get('exceptions', [])}
    tm = '2025-02-18 Text messages - Carranza to Haldane.txt'
    add('03 Readiness (pass 2: attorney selection)', 'A NO MATCH document the attorney added by hand is held for confirmation',
        f'{tm}: Names no party or key term of this matter', '; '.join(before.get(tm, [])) or 'not held',
        status(HOLD_LABELS['no_connection'] in before.get(tm, [])) if r else '',
        'The attorney confirms it belongs ("This document belongs to the matter"), which is written to the audit log, and it is released.')
    p2 = {e['name']: e['defects'] for e in (R.get('pass2') or {}).get('exceptions', [])}
    add('03 Readiness (pass 2: attorney selection)', 'Held back', 'Keystone delay notice (legacy export, Windows-1252).txt: Corrupted character encoding',
        '; '.join(f'{n}: {"; ".join(v)}' for n, v in p2.items()), status(list(p2) == ['Keystone delay notice (legacy export, Windows-1252).txt']) if r else '',
        'After the attorney confirms the text messages, only the corrupt export should be held. The scans and intruders were already excluded by screening.')
    adv = {a['label']: a for a in (R.get('pass2') or {}).get('advisories', [])}
    for label, must in [('No date found', ['Project directory - Harpeth Ridge.json', 'Punch list - Level 3 (walk with Architect).md',
                                           'Manufacturer data sheet - galvanized duct gauges.pdf']),
                        ('Very short documents', ['2025-02-18 Text messages - Carranza to Haldane.txt'])]:
        names = adv.get(label, {}).get('names', '')
        missing = [x for x in must if x not in names]
        add('03 Readiness (pass 2: attorney selection)', f'Advisory "{label}" lists', '; '.join(must),
            adv.get(label, {}).get('detail', 'not shown'), status(not missing) if r else '', 'Advisory only; never blocks.')

    # ---- Stage 04 ----
    ch = R.get('chronology') or {}
    pts = ch.get('points', [])
    by_src = Counter(p['source'] for p in pts)
    add('04 Chronology', 'Dated events assembled', 'See "Key dates" sheet for the dates that matter', ch.get('events'), 'n/a', '')
    add('04 Chronology', 'Impossible date 02/30/2025 rejected', '1 event from the Jan 15 daily report (not 2)',
        by_src.get('2025-01-15 Daily Field Report No. 196 (date typo).txt'),
        status(by_src.get('2025-01-15 Daily Field Report No. 196 (date typo).txt') == 1) if r else '', 'A rolled-over date would be a confident wrong entry.')
    ash = [p['date'] for p in pts if p['source'].startswith('2025-03-04 Ashdown')]
    add('04 Chronology', 'UK day-first date 04/03/2025 (true date March 4, 2025)', '3/4/2025', ', '.join(ash),
        ('FINDING' if ash == ['4/3/2025'] else status(ash == ['3/4/2025'])) if r else '',
        'FINDING if 4/3/2025: numeric dates are always read US month-first. A foreign sender puts the event a month off.')
    priv_on = [p for p in pts if p['source'] in {d['file'] for d in wh}]
    add('04 Chronology', 'Withheld documents excluded from the chronology', '0 events from withheld documents', len(priv_on) if r else '',
        status(not priv_on) if r else '', 'Privileged content must never reach analysis.')
    rfc_only = ['2025-02-03 Carranza to Haldane - L3 supply trunks.eml', '2024-12-19 Oyelaran to Haldane - CO 14 notice window.eml']
    add('04 Chronology', 'Email sent dates (RFC header format) on the chronology', 'Feb 3, 2025 and Dec 19, 2024 (a reviewer dates emails by when they were sent)',
        ', '.join(f'{f}: {by_src.get(f, 0)} events' for f in rfc_only), 'FINDING' if r else '',
        'FINDING: email "Date:" headers ("Mon, 3 Feb 2025") are not a format the date reader accepts, so the two hottest emails are not on the timeline at all.')

    # ---- Stage 05 ----
    cit = R.get('citations', {})
    found = missed = 0
    for d in docs:
        for p in d.get('passages', []) or []:
            if d['designation'] == 'withhold':
                continue
            hit = bool(passage_hit(p, cit.get(d['file'], [])))
            found += hit
            missed += (not hit)
    add('05 Citations', 'Must-find passages surfaced by the extractor', f'{found + missed} passages across the hot documents', f'{found} found, {missed} missed' if r else '',
        status(missed == 0, finding=True) if r else '', 'See the "Must-find passages" sheet. A miss can still be added by hand in Stage 05.')
    pdf_multi = [d['file'] for d in docs if d['fmt'] == 'pdf' and (obs_docs.get(d['file']) or {}).get('pages', 0) > 1]
    line1 = [f for f in pdf_multi if any(not c['locator'].startswith('Page') for c in cit.get(f, []))]
    if r:
        pdf_locs = '; '.join(f"{f}: {', '.join(c['locator'] for c in cit.get(f, []))}" for f in pdf_multi[:3])
        add('05 Citations', 'Citations in multi-page PDFs point to where the passage is',
            'A page for each citation', (f'{len(line1)} of {len(pdf_multi)} multi-page PDFs cite by line: ' + '; '.join(line1[:4])) if line1 else f'All cite by page. {pdf_locs}',
            status(not line1),
            'DEFECT FOUND: PDF text is extracted as one line per page, so every citation on page 1 reads "Line 1" and a passage on page 2 reads "Line 3". In a brief, "BWB-000001, Line 1" does not tell the reader where to look. Fix: cite PDFs by page.')
    hot_zero = [d['file'] for d in docs if d.get('hot') and d['designation'] != 'withhold' and d['file'] in obs_docs and not cit.get(d['file'])]
    if r:
        add('05 Citations', 'Every hot document yields at least one citation', 'All hot documents cited', '; '.join(hot_zero) or 'All cited',
            status(not hot_zero, finding=True),
            'FINDING: the extractor ranks sentences by dates, dollar amounts, full names and legal verbs. Plain-language admissions ("We ran 22 gauge", "Don\'t hold up the ceiling grid") carry none of those, so the two most damaging emails in the case get no citation at all. The attorney must add them by hand in Stage 05.')

    # ---- Stage 07-08 ----
    dl = R.get('deliverables', {})
    plog = dl.get('Privilege Log', '')
    in_log = [d['file'] for d in wh + rd if d['file'] in plog]
    add('08 Package', 'Privilege log lists every withheld and redacted document, with basis', f'{len(wh) + len(rd)} entries',
        f'{len(in_log)} of {len(wh) + len(rd)} present' if r else '', status(len(in_log) == len(wh) + len(rd)) if r else '', 'FRCP 26(b)(5)(A).')
    idx = dl.get('Production Index', '')
    intr = [d['file'] for d in docs if d.get('integrity') == 'unrelated' and d['file'] in idx]
    add('08 Package', "Other client's documents absent from the production index", 'None present', '; '.join(intr) or 'None present',
        status(not intr) if r else '', 'Confidentiality.')
    idx_rows = [ln for ln in idx.splitlines() if ln.startswith(m['batesPrefix'] + '-')]
    red_missing = [d['file'] for d in rd if d['file'] in obs_docs and not any(d['file'] in ln for ln in idx_rows)]
    if r:
        add('08 Package', 'Redacted documents listed in the production index (they are produced)', '; '.join(d['file'] for d in rd),
            ('Missing: ' + '; '.join(red_missing)) if red_missing else 'Both present', status(not red_missing),
            'DEFECT FOUND if missing: the index includes only documents marked "produce", so a document produced in redacted form carries a Bates number but appears on neither the index nor as produced. Opposing counsel receives pages the index does not account for.')
    priv_in_idx = [d['file'] for d in wh if any(d['file'] in ln for ln in idx_rows)]
    add('08 Package', 'Withheld documents not marked for production in the index', 'None', '; '.join(priv_in_idx) or 'None',
        status(not priv_in_idx) if r else '', '')
    bates = R.get('bates') or {}
    nums = list(bates.values())
    add('07 Bates', 'Bates numbers unique and prefixed BWB', f'{len(nums)} unique, all BWB-', f'{len(set(nums))} unique; prefixes {sorted({n.split("-")[0] for n in nums})}' if r else '',
        status(len(set(nums)) == len(nums) and all(n.startswith('BWB-') for n in nums)) if r else '', '')
    add('All', 'Console errors during the run', '0', len(R.get('errors', [])) if r else '', status(not R.get('errors')) if r else '', '; '.join(R.get('errors', [])[:3]))
    return C


# ---------------------------------------------------------------------------

def style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor=NAVY)
        cell.alignment = Alignment(wrap_text=True, vertical='top')


def widths(ws, ws_widths):
    for i, w in enumerate(ws_widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def write_readme(ws, m, r, checks):
    c = m['criteria']
    tally = Counter(x[4] for x in checks if x[4])
    lines = [
        ('Case Intelligence - Stress Test Answer Key', True),
        ('FICTIONAL. Every person, company, project and case number is invented for testing.', False),
        ('', False),
        (f'Matter: {m["matter"]}  (we act for Brightwater Builders, the general contractor)', True),
        ("Story: the owner (Tallis) says the Harpeth Ridge medical office building finished 74 days late and that 3rd-floor ductwork went in at the wrong gauge. It withheld $185,000 in liquidated damages. Brightwater blames Tallis's Change Order 14 (an MRI redesign), but asked for the extra time 21 days late. It blames the ductwork on its HVAC subcontractor, Keystone.", False),
        ('', False),
        (f'Collection size: {m["size"]} files (48 hand-written documents with planted facts; the rest are routine project records)', False),
        ('', False),
        ('HOW TO RUN IT BY HAND', True),
        ('1. Stage 01: enter these screening criteria, then upload every file in the folder.', False),
        (f'     Parties: {c["parties"]}', False),
        (f'     Key terms: {c["terms"]}', False),
        (f'     Period: {c["from"]} to {c["to"]}', False),
        (f'2. Stage 02: set the matter name to "{m["matter"]}" and the Bates prefix to {m["batesPrefix"]}.', False),
        ('3. Safety-net pass: select EVERY document (Select relevant, then click each unselected row). Run the Stage 03 check and compare with "Stage checks".', False),
        ('4. Attorney pass: Deselect All, Select relevant, add the text-message export by hand, remove the holiday-party email. Withhold the 4 privileged documents and redact the 2 marked Redact, using the basis and description on the "Documents" sheet. Run Stage 03 again.', False),
        ('5. Stage 04 to 09: run the analysis, check the dates on "Key dates", check the passages on "Must-find passages", approve, and download the privilege log and production index.', False),
        ('', False),
        ('HOW TO READ THE RESULTS', True),
        ('PASS = the app did what a correct review requires.  FAIL = the app got it wrong (a defect).  FINDING = the app followed its own rules, but the attorney must know the result (a limit of the approach).', False),
    ]
    if r:
        lines += [('', False), (f'AUTOMATED RUN: {r["startedAt"][:16].replace("T", " ")} UTC, {r["size"]} files, {r["totalSeconds"]:.0f} s', True),
                  (f'Results: {tally.get("PASS", 0)} PASS, {tally.get("FAIL", 0)} FAIL, {tally.get("FINDING", 0)} FINDING', True)]
    for i, (text, bold) in enumerate(lines, start=1):
        cell = ws.cell(row=i, column=1, value=text)
        cell.font = Font(bold=bold, size=14 if i == 1 else 11, color=NAVY if bold else '000000')
        cell.alignment = Alignment(wrap_text=True, vertical='top')
    ws.column_dimensions['A'].width = 150


def write_checks(ws, checks):
    head = ['Stage', 'Check', 'Expected (correct)', 'Observed (automated run)', 'Result', 'Why it matters / notes']
    ws.append(head)
    style_header(ws, 1, len(head))
    for row in checks:
        ws.append([str(x) if x is not None else '' for x in row])
        res = row[4]
        if res in FILL:
            ws.cell(row=ws.max_row, column=5).fill = PatternFill('solid', fgColor=FILL[res])
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical='top')
    widths(ws, [22, 44, 44, 44, 11, 70])
    ws.freeze_panes = 'A2'


def write_documents(ws, m, r, exp):
    head = ['#', 'File', 'Format', 'Planted?', 'True type', 'App type', 'Screening expected', 'App screening',
            'Responsive? (attorney)', 'Stage 02 call', 'Basis', 'Log description', 'Stage 03 expected',
            'Key dates (true)', 'Facts to verify', 'Attorney judgment', 'Tool note / what to watch', 'Citations found']
    ws.append(head)
    style_header(ws, 1, len(head))
    R = r or {}
    rows = R.get('rows', {})
    docs_obs = R.get('documents', {})
    order = sorted(m['documents'], key=lambda d: (not d.get('planted'), d['file']))
    for i, d in enumerate(order, start=1):
        e = exp['docs'].get(d['file'], {})
        integ = d.get('integrity')
        s3 = {'needs_ocr': 'HOLD: scanned, needs OCR', 'empty': 'HOLD: no readable text', 'corrupt': 'HOLD: corrupted encoding',
              'unrelated': 'HOLD: does not belong to this matter', 'duplicate': 'Skipped at upload (exact duplicate)'}.get(integ, 'Ready')
        if d['designation'] != 'produce':
            s3 += f' ({d["designation"]}, logged)'
        key_dates = '; '.join(f'{raw} = {iso or "NOT A VALID DATE"} ({why})' for raw, iso, why in (d.get('dates') or []))
        ws.append([i, d['file'], d['fmt'], 'Yes' if d.get('planted') else '', d['type'],
                   (docs_obs.get(d['file']) or {}).get('type', '' if not r else 'not in matter'),
                   LABEL.get(e.get('screen'), '') if integ != 'duplicate' else '', (rows.get(d['file']) or {}).get('screen', ''),
                   d.get('responsive', 'Yes'), d['designation'].capitalize(), d.get('basis', ''), d.get('description', ''), s3,
                   key_dates, d.get('facts', ''), d.get('judgment', ''), d.get('tool_note', '') or '',
                   len(R.get('citations', {}).get(d['file'], [])) if r else ''])
        if d.get('hot'):
            ws.cell(row=ws.max_row, column=2).font = Font(bold=True)
        if e and rows.get(d['file']) and rows[d['file']].get('screen') != LABEL.get(e.get('screen')):
            ws.cell(row=ws.max_row, column=8).fill = PatternFill('solid', fgColor=FILL['FAIL'])
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical='top')
    widths(ws, [5, 48, 8, 8, 11, 11, 13, 13, 12, 10, 14, 36, 22, 44, 56, 44, 56, 9])
    ws.freeze_panes = 'C2'
    ws.auto_filter.ref = ws.dimensions


def write_chronology(ws, m, r):
    head = ['True date', 'What happened', 'Source document', 'As written', 'Should be on the chronology?', 'On the app chronology?', 'Result']
    ws.append(head)
    style_header(ws, 1, len(head))
    pts = (r or {}).get('chronology', {}).get('points', [])
    have = {(p['source'], p['date']) for p in pts}
    withheld = {d['file'] for d in m['documents'] if d['designation'] == 'withhold'}
    rows = []
    for d in m['documents']:
        if not d.get('planted') or d.get('integrity') or not d.get('dates'):
            continue
        for raw, iso, why in d['dates']:
            rows.append((iso or '', why, d['file'], raw, d))
    produced = None
    if r:
        idx = r.get('deliverables', {}).get('Production Index', '')
        produced = {x['file'] for x in m['documents'] if x['file'] in idx}
        produced |= {n for n, v in (r.get('privilege') or {}).items() if v.get('status') == 'redact'}
    for iso, why, f, raw, d in sorted(rows):
        should = 'No (withheld)' if f in withheld else ('No (not a readable format)' if not iso or not dates_in(raw) else 'Yes')
        if produced is not None and f not in withheld and f not in produced:
            should = 'No (not in the production set)'
        if iso and not dates_in(raw):
            should = 'Reviewer: yes. Tool: cannot (no year / not a date format)'
        shown = ''
        res = ''
        if r:
            if iso and dates_in(raw):
                tool_d = dates_in(raw)[0][1]
                shown = 'Yes' if (f, us(tool_d)) in have else 'No'
                if f in withheld or should.startswith('No (not in'):
                    res = status(shown == 'No')
                elif tool_d.isoformat() != iso:
                    res = 'FINDING'
                    shown += f' (as {us(tool_d)})'
                else:
                    res = status(shown == 'Yes')
            elif not iso:
                bad = [p for p in pts if p['source'] == f and p['date'] in ('3/2/2025',)]
                shown = 'No' if not bad else 'Yes (rolled over)'
                res = status(not bad)
            else:
                res = 'FINDING'
                shown = 'No'
        ws.append([iso or 'invalid', why, f, raw, should, shown, res])
        if res in FILL:
            ws.cell(row=ws.max_row, column=7).fill = PatternFill('solid', fgColor=FILL[res])
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical='top')
    widths(ws, [12, 48, 52, 18, 26, 18, 10])
    ws.freeze_panes = 'A2'


def write_passages(ws, m, r):
    head = ['Document', 'Passage an attorney would cite', 'Why it matters', 'Extractor surfaced it?', 'Result']
    ws.append(head)
    style_header(ws, 1, len(head))
    cit = (r or {}).get('citations', {})
    for d in m['documents']:
        for p in d.get('passages', []) or []:
            if d['designation'] == 'withhold':
                continue
            hit = passage_hit(p, cit.get(d['file'], []))
            res = ('PASS' if hit else 'FINDING') if r else ''
            ws.append([d['file'], p, d.get('facts', ''), (hit[0]['locator'] if hit else 'No - add by hand in Stage 05') if r else '', res])
            if res:
                ws.cell(row=ws.max_row, column=5).fill = PatternFill('solid', fgColor=FILL[res])
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical='top')
    widths(ws, [46, 60, 60, 22, 10])


def write_scale(ws, dirs):
    head = ['Files', 'Accepted', 'Upload + index (s)', 'Stage 02 list (s)', 'Readiness check, all (s)', 'Readiness check, attorney set (s)',
            'Analysis (s)', 'Total run (s)', 'Browser storage (MB)', 'JS heap (MB)', 'Pass 1 state', 'Pass 1 held',
            'Intruders caught (of 2)', 'Chronology events', 'Console errors']
    ws.append(head)
    style_header(ws, 1, len(head))
    for d in dirs:
        m, r = load(d)
        if not r:
            ws.append([m['size'], 'not run'])
            continue
        t = r['timings']
        held = [e['name'] for e in r['pass1']['exceptions']]
        intr = [x['file'] for x in m['documents'] if x.get('integrity') == 'unrelated' and x['file'] in held]
        ws.append([m['size'], r.get('accepted'), t.get('upload_and_index'), t.get('stage02_render_rows'), t.get('pass1_readiness_check'),
                   t.get('pass2_readiness_check'), t.get('stage04_analysis'), round(r['totalSeconds']), r['storage']['usageMB'],
                   r['storage']['heapMB'], r['pass1']['state'], len(held), len(intr), r['chronology']['events'], len(r['errors'])])
    widths(ws, [8, 9, 12, 12, 14, 14, 11, 11, 12, 11, 16, 10, 12, 12, 10])


if __name__ == '__main__':
    ap = argparse.ArgumentParser()
    ap.add_argument('dir')
    ap.add_argument('--scale', nargs='*', default=[])
    a = ap.parse_args()
    out, checks = build_workbook(a.dir, a.scale)
    t = Counter(c[4] for c in checks if c[4])
    print(out)
    print(dict(t))
